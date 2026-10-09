"""R07 — contrato de reimportação segura, somente com planilhas de demonstração.

Os cenários identificados na base antiga ficam @expectedFailure até o Claude
entregar a correção em legado.py. Remover a anotação quando passarem.
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.worksheet.table import Table

from ascalpi_producao import legado
from ascalpi_producao.servico import Servico
from tests.apoio import criar_controle, criar_livro


class TestReimportacaoSegura(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        raiz = Path(self.tmp.name)
        self.origem = raiz / "origem"
        self.origem.mkdir()
        criar_livro(self.origem)
        self.controle = criar_controle(self.origem)
        self.servico = Servico(raiz / "dados")
        self.servico.salvar_config({"publicar": False})
        legado.importar_pasta(self.servico.banco, self.origem, raiz / "dados")

    def tearDown(self):
        self.servico.banco.fechar()
        self.tmp.cleanup()

    def _legado(self):
        return self.servico.banco.todos("SELECT id, numero, origem, status_instalacao, obs FROM ops WHERE origem = 'LEGADO' ORDER BY id")

    @unittest.expectedFailure
    def test_r07_reimportacao_idempotente_preserva_identidade_e_anotacoes(self):
        """Uma nova importação do mesmo arquivo não deve apagar edição do acompanhamento."""
        antes = self._legado()
        self.assertEqual(len(antes), 2)
        alvo = antes[0]["id"]
        self.servico.acompanhar(alvo, {"obs": "OBSERVAÇÃO REGISTRADA NO SISTEMA", "status_instalacao": "OK"})
        legado.importar_controle(self.servico.banco, self.controle, [])
        depois = self._legado()
        self.assertEqual({x["id"] for x in antes}, {x["id"] for x in depois}, "IDS DO HISTÓRICO MUDARAM")
        registro = next(x for x in depois if x["id"] == alvo)
        self.assertEqual(registro["obs"], "OBSERVAÇÃO REGISTRADA NO SISTEMA")
        self.assertEqual(registro["status_instalacao"], "OK")

    @unittest.expectedFailure
    def test_r07_tabela_ausente_nao_remove_registros(self):
        """Arquivo sem Controle_OP: rejeição explícita, sem mutação."""
        antes = self._legado()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["B2"] = "Nº O.P"
        ws["B3"] = "888-26"
        ws.add_table(Table(displayName="Outra_Tabela", ref="B2:B3"))
        invalido = self.origem / "controle_invalido.xlsx"
        wb.save(invalido)
        with self.assertRaises(ValueError):
            legado.importar_controle(self.servico.banco, invalido, [])
        self.assertEqual(self._legado(), antes)

    @unittest.expectedFailure
    def test_r07_tabela_sem_coluna_numero_nao_remove_registros(self):
        """Tabela chamada Controle_OP, mas sem coluna obrigatória: bloquear."""
        antes = self._legado()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["B2"] = "OBRA"
        ws["B3"] = "PRAÇA DO TESTE"
        ws.add_table(Table(displayName="Controle_OP", ref="B2:B3"))
        invalido = self.origem / "controle_sem_numero.xlsx"
        wb.save(invalido)
        with self.assertRaises(ValueError):
            legado.importar_controle(self.servico.banco, invalido, [])
        self.assertEqual(self._legado(), antes)


if __name__ == "__main__":
    unittest.main()
