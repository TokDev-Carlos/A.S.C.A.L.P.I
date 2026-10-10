import io
import re
import tempfile
import unittest
import zipfile
from datetime import date, datetime
from pathlib import Path

import openpyxl

from ascalpi_producao import documento as d, pacote
from tests.apoio import criar_livro


class TestDocumento(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.livro = criar_livro(Path(cls.tmp.name))
        cls.modelo = pacote.extrair_aba(cls.livro, "O.P-ATA-TESTE")

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_extrair_so_a_aba(self):
        z = zipfile.ZipFile(io.BytesIO(self.modelo))
        self.assertEqual([a["nome"] for a in pacote.abas(pacote.carregar(io.BytesIO(self.modelo)))], ["O.P-ATA-TESTE"])
        self.assertNotIn("xl/vbaProject.bin", z.namelist())
        self.assertIn('fullCalcOnLoad="1"', z.read("xl/workbook.xml").decode())
        self.assertEqual(d.cabecalho_do_modelo(self.modelo)["cliente"], "CIDADE TESTE-RJ")
        self.assertEqual([l["codigo"] for l in d.linhas_do_modelo(self.modelo)], ["1", "2", "2.1", "3", "0.1"])

    def test_plano_paginas(self):
        self.assertEqual(d.plano_paginas(5), [5])
        self.assertEqual(d.plano_paginas(13), [13])
        self.assertEqual(d.plano_paginas(14), [13, 1])
        self.assertEqual(d.plano_paginas(29), [13, 8, 8])
        self.assertEqual(d.plano_paginas(43), [13, 15, 15])

    def test_textos(self):
        self.assertEqual(d.texto_dia_prazo(date(2026, 10, 30)), "Sexta-Feira" + " " * 48 + "30 Outubro 2026")
        self.assertEqual(d.texto_dia_prazo("DEFINIR"), "Definir")
        self.assertEqual(d.texto_hoje(datetime(2026, 10, 9, 8, 5)), "Hoje 09/10/2026 8:05-sexta-feira")

    def _gerar(self, prazo, linhas, paginacao="altura"):
        dados = d.DadosOP("010-26", "PRAÇA <&> TESTE", prazo, "ANA", linhas, rev=2, emitido_em=datetime(2026, 10, 9, 8, 5))
        return openpyxl.load_workbook(io.BytesIO(d.gerar_xlsx(self.modelo, dados, "senha", paginacao)))

    def test_gerar_xlsx(self):
        wb = self._gerar(date(2026, 10, 30), {10: d.LinhaOP(2, None, "AZUL"), 13: d.LinhaOP(1.5, "X")})
        ws = wb.active
        self.assertEqual(ws["P2"].value, "010-26")
        self.assertEqual(ws["U2"].value, 2)
        self.assertEqual(ws["D3"].value, "PRAÇA <&> TESTE")
        self.assertEqual(ws["V2"].value, d.serial_excel(date(2026, 10, 30)))  # célula com formato de data no modelo real
        self.assertEqual(ws["V8"].value, "ANA")
        self.assertEqual(ws["D10"].value, 2)
        self.assertEqual(ws["S10"].value, "AZUL")
        self.assertEqual(ws["D13"].value, 1.5)
        self.assertEqual(ws["E13"].value, "X")
        self.assertIsNone(ws["D12"].value)                       # resto do modelo limpo
        self.assertFalse(ws.row_dimensions[10].hidden)
        self.assertTrue(ws.row_dimensions[11].hidden)
        self.assertTrue(ws.row_dimensions[100].hidden)
        self.assertEqual(ws.print_area, "'O.P-ATA-TESTE'!$B$2:$V$13")
        self.assertEqual(ws.page_setup.orientation, "portrait")
        self.assertEqual(ws["B13"].border.bottom.style, "thin")
        # somente visualizar
        self.assertTrue(ws.protection.sheet)
        self.assertEqual(ws.protection.algorithmName, "SHA-512")
        self.assertTrue(wb.security.lockStructure)
        self.assertFalse(any(c.protection.locked is False for r in ws.iter_rows() for c in r))

    def test_prazo_texto_e_quebras(self):
        linhas = {r: d.LinhaOP(1) for r in range(10, 40)}
        wb = self._gerar("DEFINIR", linhas, paginacao="legado")
        ws = wb.active
        self.assertEqual(ws["V2"].value, "DEFINIR")
        self.assertEqual(ws["S6"].value, "Definir")
        self.assertEqual([b.id for b in ws.row_breaks.brk], [22, 31])   # motor legado (VBA): 13 + 9 + 8

    def test_sem_itens(self):
        with self.assertRaises(ValueError):
            d.gerar_xlsx(self.modelo, d.DadosOP("1", "O", "DEFINIR", "A", {}), "x")

    def test_hash_senha(self):
        h = d.hash_senha("teste", b"\x00" * 16, giros=1)
        self.assertEqual(h["algorithmName"], "SHA-512")
        self.assertTrue(re.fullmatch(r"[A-Za-z0-9+/=]{88}", h["hashValue"]))


if __name__ == "__main__":
    unittest.main()
