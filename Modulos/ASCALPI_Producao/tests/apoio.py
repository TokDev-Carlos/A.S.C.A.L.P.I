"""Livro legado sintético (sem dados de produção) para os testes."""
from pathlib import Path

import openpyxl
from openpyxl.worksheet.table import Table

EQUIPAMENTOS = [("1", "BALANÇO INOX"), ("2", "GANGORRA"), ("2.1", "GANGORRA - AZUL"), ("3", "ESCORREGA"), ("0.1", "TOTEM BONIFICADO")]


def criar_livro(pasta: Path, nome: str = "OK-TESTE_O.P.xlsx", cliente: str = "CIDADE TESTE-RJ") -> Path:
    wb = openpyxl.Workbook()
    saldo = wb.active
    saldo.title = "CONTRATO TESTE"
    saldo.sheet_properties.codeName = "Saldo_TESTE"
    saldo["E2"] = "Contrato: \nFOR-01/2026"
    cab = ["COD", "IMG", "EQUIPAMENTOS", "VALOR UN.", "MONTANTE", "1ª", "ACUMULADO", "MEDINDO", "QUANT.", "SALDO", "O.P 001-26", "PREVISÃO"]
    for i, h in enumerate(cab, 1):
        saldo.cell(3, i, h)
    linhas = [(1, "BALANÇO INOX COMPLETO", 10, 2, 3), (2, "GANGORRA DUPLA", 5, 1, 4), (3, "ESCORREGA", 0, 0, 0), ("0.1", "TOTEM", 0, 0, 0)]
    for r, (cod, eq, mont, quant, prev) in enumerate(linhas, 4):
        saldo.cell(r, 1, cod); saldo.cell(r, 3, eq); saldo.cell(r, 4, 100.0); saldo.cell(r, 5, mont)
        saldo.cell(r, 9, quant); saldo.cell(r, 12, prev)
    saldo.cell(8, 1, "TOTAL DO CONTRATO:")
    saldo.add_table(Table(displayName="TAB_SALDO_TESTE", ref="A3:L8"))

    op = wb.create_sheet("O.P-ATA-TESTE")
    op.sheet_properties.codeName = "OP_TESTE"
    op["B2"] = '="Hoje "&TEXT(NOW(),"dd/mm/aaaa")'
    op["D2"] = "ORDEM DE PRODUÇÃO"; op["L2"] = "AÇO"; op["N2"] = "OP:"; op["P2"] = "999-99"; op["T2"] = "REV:"; op["U2"] = 0
    op["B3"] = cliente; op["S3"] = "EQUIPAMENTOS MODELO:"; op["S4"] = "ATA TESTE"; op["B8"] = "MOB"
    op["S6"] = '=PROPER(TEXT(V2,"dddd"))'
    for faixa in ("B3:B4", "D3:R4", "P2:S2", "S6:V7"):
        op.merge_cells(faixa)
    for r, (cod, equipamento) in enumerate(EQUIPAMENTOS, 10):
        op.cell(r, 1, cod); op.cell(r, 2, equipamento)
    for r in range(10, 101):
        op.merge_cells(f"E{r}:F{r}"); op.merge_cells(f"S{r}:V{r}")
    op["D12"] = 7  # resto de uso anterior: deve ser limpo
    op.protection.sheet = True

    copia = wb.copy_worksheet(op)          # cópia com o padrão OP- : modelo válido, mesmo contrato
    copia.title = "O.P-ATA-TESTE-COPIA"
    copia.sheet_properties.codeName = "OP_TESTE2"
    solta = wb.copy_worksheet(op)          # cópia sem o padrão no nome: não é modelo
    solta.title = "valadão"
    solta.sheet_properties.codeName = "OP_TESTE1"

    db = wb.create_sheet("DB")
    db.sheet_properties.codeName = "DB_"
    db["B2"] = "CLIENTE"; db["B3"] = cliente
    db.add_table(Table(displayName="TAB_CLIENTE", ref="B2:B3"))
    db.sheet_state = "veryHidden"
    destino = Path(pasta) / nome
    wb.save(destino)
    return destino


def criar_controle(pasta: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CONTROLE_ORDENS_DE_PRODUÇÃO"
    cab = ["Nº\nO.P", "CLIENTE", "OBRA", "SOLICITANTE", "EQUIPAMENTOS", "TIPO DE MATERIAL", "SOLICITAÇÃO",
           "DATA ENTREGA", "DATA ENTREGA\nATUALIZADA", "MATERIAL OBRA", "STATUS INSTALAÇÃO", "FOTOGRÁFICO", "OBS:"]
    for i, h in enumerate(cab, 2):
        ws.cell(2, i, h)
    import datetime as dt
    ws.append([None, "005-26", "CIDADE TESTE-RJ", "PRAÇA A", "ANA", "MOB", "CARBONO", dt.datetime(2026, 1, 5, 9), dt.datetime(2026, 1, 20), None, None, "OK"])
    ws.append([None, "012-26", "CIDADE TESTE", "PRAÇA B", "LUIS", "MOB", "INOX", dt.datetime(2026, 2, 1, 9), "DEFINIR"])
    ws.add_table(Table(displayName="Controle_OP", ref="B2:N4"))
    destino = Path(pasta) / "2_Controle_Ordens_De_Producao.xlsx"
    wb.save(destino)
    return destino
