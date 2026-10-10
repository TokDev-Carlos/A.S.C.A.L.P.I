"""G1 — edição do modelo no Excel: cópia de trabalho, validação, nova versão e histórico imutável."""
import io
import json
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

import openpyxl

from ascalpi_producao import legado
from ascalpi_producao.modelos_edicao import ErroConflitoEdicao
from ascalpi_producao.servico import ErroValidacao, Servico
from ascalpi_producao.servidor import servir
from tests.apoio import criar_controle, criar_livro


class TestEdicaoModelo(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.base = Path(self.tmp.name)
        (self.base / "legado").mkdir()
        criar_livro(self.base / "legado")
        criar_controle(self.base / "legado")
        self.s = Servico(self.base / "dados")
        self.s.salvar_config({"publicar": False, "edicao_modelo_excel": True})
        legado.importar_pasta(self.s.banco, self.base / "legado", self.base / "dados")
        self.modelo = next(m for m in self.s.modelos() if m["aba"] == "O.P-ATA-TESTE")

    def tearDown(self):
        self.s.banco.fechar()
        self.tmp.cleanup()

    # ------------------------------------------------------------ apoio
    def _copia(self, edicao: dict) -> Path:
        e = self.s.banco.um("SELECT arquivo_trabalho FROM modelo_edicoes WHERE id = ?", (edicao["id"],))
        return self.s.dados / e["arquivo_trabalho"]

    def _editar(self, edicao: dict, mudar) -> None:
        caminho = self._copia(edicao)
        wb = openpyxl.load_workbook(caminho)
        mudar(wb.active)
        wb.save(caminho)

    def _op(self, linha_cod="1"):
        linhas = {l["codigo"]: l["linha"] for l in self.s.modelo_completo(self.modelo["id"])["linhas"]}
        return self.s.salvar_op({"modelo_id": self.modelo["id"], "obra": "praça", "solicitante": "ana", "prazo": "definir",
                                 "itens": [{"linha": linhas[linha_cod], "quantidade": "1"}]})

    # ------------------------------------------------------------ testes
    def test_desligada_por_padrao(self):
        self.s.salvar_config({"edicao_modelo_excel": False})
        with self.assertRaises(ErroValidacao):
            self.s.edicao.iniciar(self.modelo["id"])
        with self.assertRaises(ErroValidacao):
            self.s.salvar_config({"edicao_modelo_excel": "true"})   # texto não liga a função

    def test_copia_editavel_e_unica(self):
        e1 = self.s.edicao.iniciar(self.modelo["id"])
        e2 = self.s.edicao.iniciar(self.modelo["id"])
        self.assertEqual(e1["id"], e2["id"])
        with zipfile.ZipFile(self._copia(e1)) as z:
            xml = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("xl/worksheets/sheet"))
        self.assertNotIn("<sheetProtection", xml)
        publicado = self.s.dados / self.modelo_arquivo()
        self.assertIn(b"sheetProtection", zipfile.ZipFile(publicado).read(
            [n for n in zipfile.ZipFile(publicado).namelist() if n.startswith("xl/worksheets/sheet")][0]))

    def modelo_arquivo(self):
        return self.s.banco.um("SELECT arquivo FROM modelos WHERE id = ?", (self.modelo["id"],))["arquivo"]

    def test_sem_alteracao_nao_publica(self):
        e = self.s.edicao.iniciar(self.modelo["id"])
        r = self.s.edicao.validar(e["id"])
        self.assertTrue(r["ok"])
        self.assertEqual(r["diferencas"], [])
        with self.assertRaises(ErroValidacao):
            self.s.edicao.publicar(e["id"], "teste", r["hash"])

    def test_publicar_nova_versao_preserva_historico(self):
        op_antiga = self._op("1")["op"]
        arquivo_antigo = self.modelo_arquivo()
        e = self.s.edicao.iniciar(self.modelo["id"])

        def mudar(ws):
            ws["B10"] = "BALANÇO INOX NOVO"
            ws.row_dimensions[11].height = 120
            ws["A20"], ws["B20"] = "9", "PEÇA SEM CONTRATO"
        self._editar(e, mudar)
        r = self.s.edicao.validar(e["id"])
        self.assertTrue(r["ok"], r["erros"])
        tipos = {(d["tipo"], d["linha"]) for d in r["diferencas"]}
        self.assertIn(("NOME", 10), tipos)
        self.assertIn(("ALTURA", 11), tipos)
        self.assertIn(("ADICIONADO", 20), tipos)
        self.assertTrue(any("SEM ITEM NO CONTRATO" in a and "9" in a for a in r["avisos"]))

        with self.assertRaises(ValueError):                         # motivo obrigatório
            self.s.edicao.publicar(e["id"], "", r["hash"])
        with self.assertRaises(ErroConflitoEdicao):                 # hash diferente do validado
            self.s.edicao.publicar(e["id"], "foto nova", "0" * 64)
        res = self.s.edicao.publicar(e["id"], "nome novo do balanço", r["hash"])
        self.assertTrue(res["ok"])
        novo = self.modelo_arquivo()
        self.assertNotEqual(novo, arquivo_antigo)
        self.assertTrue((self.s.dados / arquivo_antigo).exists())   # versão anterior nunca é apagada nem alterada
        linhas = {l["linha"]: l for l in self.s.modelo_completo(self.modelo["id"])["linhas"]}
        self.assertEqual(linhas[10]["equipamento"], "BALANÇO INOX NOVO")
        self.assertIn(20, linhas)
        # O.P. emitida antes continua no modelo da emissão
        o = self.s.banco.um("SELECT modelo_arquivo FROM ops WHERE id = ?", (op_antiga["id"],))
        self.assertEqual(o["modelo_arquivo"], arquivo_antigo)
        nome, xlsx = self.s.gerar_documento(op_antiga["id"], "xlsx")
        ws = openpyxl.load_workbook(io.BytesIO(xlsx)).active
        self.assertEqual(ws["B10"].value, "BALANÇO INOX")
        # O.P. nova usa a versão nova
        nova = self._op("1")["op"]
        self.assertEqual(self.s.banco.um("SELECT modelo_arquivo FROM ops WHERE id = ?", (nova["id"],))["modelo_arquivo"], novo)
        # edição finalizada não pode ser publicada de novo
        with self.assertRaises(ErroConflitoEdicao):
            self.s.edicao.publicar(e["id"], "de novo", r["hash"])
        eventos = [ev["acao"] for ev in self.s.eventos()]
        self.assertIn("MODELO_VERSAO_PUBLICADA", eventos)

    def test_conflito_se_modelo_mudou_durante_a_edicao(self):
        e = self.s.edicao.iniciar(self.modelo["id"])
        self._editar(e, lambda ws: ws.__setitem__("B10", "OUTRO NOME"))
        r = self.s.edicao.validar(e["id"])
        antes = set((self.s.dados / "Modelos").rglob("*.xlsx"))
        with self.s.banco.transacao() as con:     # outra versão publicada por outro caminho
            con.execute("UPDATE modelos SET arquivo = ? WHERE id = ?", (self.modelo_arquivo() + ".x", self.modelo["id"]))
        with self.assertRaises(ErroConflitoEdicao):
            self.s.edicao.publicar(e["id"], "motivo", r["hash"])
        self.assertEqual(set((self.s.dados / "Modelos").rglob("*.xlsx")), antes)   # nenhum arquivo órfão

    def test_arquivo_invalido_e_codigo_repetido(self):
        e = self.s.edicao.iniciar(self.modelo["id"])
        self._editar(e, lambda ws: ws.__setitem__("A11", "1"))      # código 1 repetido (linha 10 já é 1)
        r = self.s.edicao.validar(e["id"])
        self.assertFalse(r["ok"])
        self.assertTrue(any("REPETIDO" in x for x in r["erros"]))
        caminho = self._copia(e)
        wb = openpyxl.load_workbook(caminho)
        wb.create_sheet("EXTRA")
        wb.save(caminho)
        self.assertTrue(any("1 ABA" in x for x in self.s.edicao.validar(e["id"])["erros"]))
        caminho.write_bytes(b"isto nao e um xlsx")
        r = self.s.edicao.validar(e["id"])
        self.assertFalse(r["ok"])
        with self.assertRaises(ErroValidacao):
            self.s.edicao.publicar(e["id"], "motivo", r["hash"])

    def test_descartar_guarda_copia(self):
        e = self.s.edicao.iniciar(self.modelo["id"])
        d = self.s.edicao.descartar(e["id"])
        self.assertEqual(d["estado"], "DESCARTADA")
        self.assertTrue(self._copia(e).exists())
        with self.assertRaises(ErroConflitoEdicao):
            self.s.edicao.validar(e["id"])
        self.assertNotEqual(self.s.edicao.iniciar(self.modelo["id"])["id"], e["id"])

    def test_abrir_fora_do_windows(self):
        e = self.s.edicao.iniciar(self.modelo["id"])
        import os
        if os.name != "nt":
            with self.assertRaises(ErroValidacao):
                self.s.edicao.abrir(e["id"])

    def test_reimportar_livro_mantem_modelo_editado(self):
        e = self.s.edicao.iniciar(self.modelo["id"])
        self._editar(e, lambda ws: ws.__setitem__("B10", "EDITADO NO SISTEMA"))
        r = self.s.edicao.validar(e["id"])
        self.s.edicao.publicar(e["id"], "ajuste", r["hash"])
        publicado = self.modelo_arquivo()
        relatorio = legado.importar_pasta(self.s.banco, self.base / "legado", self.base / "dados")
        self.assertEqual(self.modelo_arquivo(), publicado)
        self.assertTrue(any("EDITADO NO SISTEMA" in x for x in relatorio))
        linhas = {l["linha"]: l for l in self.s.modelo_completo(self.modelo["id"])["linhas"]}
        self.assertEqual(linhas[10]["equipamento"], "EDITADO NO SISTEMA")

    def test_api_exige_cabecalho_da_tela(self):
        httpd = servir(self.s, porta=0)
        t = threading.Thread(target=httpd.serve_forever, daemon=True)
        t.start()
        try:
            url = f"http://127.0.0.1:{httpd.server_address[1]}/api/modelos/{self.modelo['id']}/edicao"

            def post(cabecalho):
                req = urllib.request.Request(url, data=b"{}", method="POST",
                                             headers={"Content-Type": "application/json", **cabecalho})
                try:
                    with urllib.request.urlopen(req) as r:
                        return r.status, json.loads(r.read())
                except urllib.error.HTTPError as erro:
                    return erro.code, json.loads(erro.read())
            self.assertEqual(post({})[0], 403)
            status, corpo = post({"X-ASCALPI": "1"})
            self.assertEqual(status, 200)
            self.assertEqual(corpo["estado"], "ABERTA")
            with urllib.request.urlopen(url) as r:
                self.assertEqual(json.loads(r.read())["edicao"]["id"], corpo["id"])
        finally:
            httpd.shutdown()
            httpd.server_close()

    # ------------------------------------------------------------ revisão própria do G1 (10/10/2026)
    def test_config_em_texto_nao_liga_a_funcao(self):
        cfg = json.loads(self.s.arquivo_config.read_text(encoding="utf-8"))
        cfg["edicao_modelo_excel"] = "false"                        # config.json editado à mão
        self.s.arquivo_config.write_text(json.dumps(cfg), encoding="utf-8")
        with self.assertRaises(ErroValidacao):
            self.s.edicao.iniciar(self.modelo["id"])

    def test_iniciar_concorrente_devolve_a_mesma_edicao(self):
        resultados, falhas = [], []
        barreira = threading.Barrier(4)

        def iniciar():
            barreira.wait()
            try:
                resultados.append(self.s.edicao.iniciar(self.modelo["id"])["id"])
            except Exception as erro:                               # pragma: no cover - falha do teste
                falhas.append(erro)
        ts = [threading.Thread(target=iniciar) for _ in range(4)]
        for t in ts:
            t.start()
        for t in ts:
            t.join()
        self.assertEqual(falhas, [])
        self.assertEqual(len(set(resultados)), 1)
        self.assertEqual(self.s.banco.um("SELECT COUNT(*) n FROM modelo_edicoes")["n"], 1)

    def test_origem_apagada_vira_conflito(self):
        e = self.s.edicao.iniciar(self.modelo["id"])
        self._editar(e, lambda ws: ws.__setitem__("B10", "OUTRO NOME"))
        r = self.s.edicao.validar(e["id"])
        origem = self.modelo_arquivo()
        with self.s.banco.transacao() as con:     # reimportação trocou o arquivo e apagou o antigo (sem uso)
            con.execute("UPDATE modelos SET arquivo = ? WHERE id = ?", (origem + ".novo.xlsx", self.modelo["id"]))
        (self.s.dados / origem).unlink()
        with self.assertRaises(ErroConflitoEdicao):
            self.s.edicao.validar(e["id"])
        with self.assertRaises(ErroConflitoEdicao):
            self.s.edicao.publicar(e["id"], "motivo", r["hash"])

    def test_publicar_exige_a_validacao_gravada(self):
        e = self.s.edicao.iniciar(self.modelo["id"])
        self._editar(e, lambda ws: ws.__setitem__("B10", "SEM VALIDAR"))
        import hashlib
        h = hashlib.sha256(self._copia(e).read_bytes()).hexdigest()   # hash certo, mas nunca validado
        with self.assertRaises(ErroConflitoEdicao):
            self.s.edicao.publicar(e["id"], "motivo", h)
        r = self.s.edicao.validar(e["id"])
        self.assertEqual(r["hash"], h)
        self.assertTrue(self.s.edicao.publicar(e["id"], "motivo", h)["ok"])

    def test_ata_s4_e_tipo_publicados_sobrevivem_a_reimportacao(self):
        e = self.s.edicao.iniciar(self.modelo["id"])

        def mudar(ws):
            ws["S4"] = "ATA NOVA 2027"
            ws["B8"] = "TIPO NOVO"
        self._editar(e, mudar)
        r = self.s.edicao.validar(e["id"])
        self.assertIn("ATA (S4)", {d.get("campo") for d in r["diferencas"]})
        self.s.edicao.publicar(e["id"], "ata nova", r["hash"])
        m = self.s.banco.um("SELECT titulo, tipo_padrao, aba FROM modelos WHERE id = ?", (self.modelo["id"],))
        self.assertEqual((m["titulo"], m["tipo_padrao"]), ("ATA NOVA 2027", "TIPO NOVO"))
        legado.importar_pasta(self.s.banco, self.base / "legado", self.base / "dados")
        depois = self.s.banco.um("SELECT titulo, tipo_padrao, aba FROM modelos WHERE id = ?", (self.modelo["id"],))
        self.assertEqual(dict(depois), dict(m))

    def test_arquivo_publicado_sumiu_reimportacao_desmarca_editado(self):
        e = self.s.edicao.iniciar(self.modelo["id"])
        self._editar(e, lambda ws: ws.__setitem__("B10", "EDITADO"))
        r = self.s.edicao.validar(e["id"])
        self.s.edicao.publicar(e["id"], "ajuste", r["hash"])
        (self.s.dados / self.modelo_arquivo()).unlink()            # versão do sistema perdida
        legado.importar_pasta(self.s.banco, self.base / "legado", self.base / "dados")
        m = self.s.banco.um("SELECT editado_sistema FROM modelos WHERE id = ?", (self.modelo["id"],))
        self.assertEqual(m["editado_sistema"], 0)                  # o livro volta a valer nas próximas importações

    def test_aba_renomeada(self):
        e = self.s.edicao.iniciar(self.modelo["id"])
        self._editar(e, lambda ws: setattr(ws, "title", "PLANILHA1"))
        r = self.s.edicao.validar(e["id"])
        self.assertFalse(r["ok"])
        self.assertTrue(any("OP-" in x for x in r["erros"]))
        self._editar(e, lambda ws: setattr(ws, "title", "O.P-ATA-TESTE2"))
        r = self.s.edicao.validar(e["id"])
        self.assertTrue(r["ok"], r["erros"])
        self.assertIn("ABA", {d.get("campo") for d in r["diferencas"]})
        self.s.edicao.publicar(e["id"], "aba nova", r["hash"])
        self.assertEqual(self.s.banco.um("SELECT aba FROM modelos WHERE id = ?", (self.modelo["id"],))["aba"],
                         "O.P-ATA-TESTE2")


    # ------------------------------------------------------------ V-G1 (Codex, 10/10/2026)
    def _publicar(self, texto: str) -> str:
        e = self.s.edicao.iniciar(self.modelo["id"])
        self._editar(e, lambda ws: ws.__setitem__("B10", texto))
        r = self.s.edicao.validar(e["id"])
        self.s.edicao.publicar(e["id"], texto, r["hash"])
        return self.modelo_arquivo()

    def test_vg1_01_reimportacao_preserva_todas_as_versoes_publicadas(self):
        original = self.modelo_arquivo()
        v1 = self._publicar("VERSÃO 1")
        v2 = self._publicar("VERSÃO 2")                            # nenhuma O.P. usa v1 nem o original
        legado.importar_pasta(self.s.banco, self.base / "legado", self.base / "dados")
        for arq in (original, v1, v2):
            self.assertTrue((self.s.dados / arq).exists(), arq)
        historico = self.s.banco.todos("SELECT arquivo_origem, arquivo_publicado FROM modelo_edicoes "
                                       "WHERE estado = 'PUBLICADA' ORDER BY id")
        self.assertEqual([(h["arquivo_origem"], h["arquivo_publicado"]) for h in historico], [(original, v1), (v1, v2)])
        for h in historico:
            for arq in (h["arquivo_origem"], h["arquivo_publicado"]):
                self.assertTrue((self.s.dados / arq).is_file(), arq)
        self.assertEqual(self.modelo_arquivo(), v2)

    def test_vg1_02_origem_alterada_no_mesmo_caminho_e_conflito(self):
        e = self.s.edicao.iniciar(self.modelo["id"])
        self._editar(e, lambda ws: ws.__setitem__("B10", "EDITADO NA CÓPIA"))
        r = self.s.edicao.validar(e["id"])                          # validação boa, antes da mudança externa
        origem = self.s.dados / self.modelo_arquivo()
        wb = openpyxl.load_workbook(origem)
        wb.active["B11"] = "MUDADO POR FORA"
        wb.save(origem)                                            # mesmo caminho, conteúdo diferente
        antes_banco = (self.modelo_arquivo(),
                       self.s.banco.um("SELECT estado FROM modelo_edicoes WHERE id = ?", (e["id"],))["estado"])
        antes_arquivos = set((self.s.dados / "Modelos").rglob("*.xlsx"))
        with self.assertRaises(ErroConflitoEdicao):
            self.s.edicao.validar(e["id"])
        with self.assertRaises(ErroConflitoEdicao):
            self.s.edicao.publicar(e["id"], "motivo", r["hash"])
        self.assertEqual((self.modelo_arquivo(),
                          self.s.banco.um("SELECT estado FROM modelo_edicoes WHERE id = ?", (e["id"],))["estado"]),
                         antes_banco)
        self.assertEqual(set((self.s.dados / "Modelos").rglob("*.xlsx")), antes_arquivos)   # nenhum arquivo órfão


class TestProtecaoDaApi(unittest.TestCase):
    """Toda ação que grava exige o cabeçalho da tela; nomes de domínio no Host são recusados (DNS rebinding)."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.s = Servico(Path(self.tmp.name) / "dados")
        self.httpd = servir(self.s, porta=0)
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()
        self.base = f"http://127.0.0.1:{self.httpd.server_address[1]}"

    def tearDown(self):
        self.httpd.shutdown()
        self.httpd.server_close()
        self.s.banco.fechar()
        self.tmp.cleanup()

    def _pedir(self, metodo, url, cabecalhos):
        req = urllib.request.Request(self.base + url, data=b"{}" if metodo != "GET" else None, method=metodo,
                                     headers=cabecalhos)
        try:
            with urllib.request.urlopen(req) as r:
                return r.status
        except urllib.error.HTTPError as erro:
            return erro.code

    def test_post_sem_cabecalho_e_recusado(self):
        for metodo, url in (("POST", "/api/ops"), ("POST", "/api/ops/1/cancelar"), ("PUT", "/api/config")):
            self.assertEqual(self._pedir(metodo, url, {"Content-Type": "text/plain"}), 403, url)

    def test_host_com_nome_de_dominio_e_recusado(self):
        self.assertEqual(self._pedir("GET", "/api/painel", {"Host": "ataque.example:8765"}), 403)
        self.assertEqual(self._pedir("GET", "/api/painel", {"Host": "localhost"}), 200)
        self.assertEqual(self._pedir("GET", "/api/painel", {}), 200)


if __name__ == "__main__":
    unittest.main()
