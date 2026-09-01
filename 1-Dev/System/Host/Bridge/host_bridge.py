from __future__ import annotations

import base64
import hashlib
import hmac
import json
import os
import secrets
import sqlite3
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo


def _add_runtime_site_packages(root: Path) -> None:
    site_packages = root / "Runtime" / "Python" / "Lib" / "site-packages"
    if site_packages.is_dir():
        value = str(site_packages)
        if value not in sys.path:
            sys.path.insert(0, value)


def _master(value: str) -> Path:
    root = Path(value).expanduser().resolve()
    app = root / "App"
    _add_runtime_site_packages(root)
    if not (app / "Config" / "master.id").is_file():
        raise RuntimeError("MESTRE CJL System LAYOUT V4 INVALIDO.")
    if str(app) not in sys.path:
        sys.path.insert(0, str(app))
    os.environ["CJL_NETWORK_ROOT"] = str(root)
    return root


def _configure_install(application_app: str, master: str, install_root: str) -> tuple[Path, Path]:
    app = Path(application_app).resolve()
    project = app.parent
    master_root = Path(master).resolve()
    _add_runtime_site_packages(Path(install_root).resolve())
    if not (app / "Config" / "sistema.json").is_file():
        raise RuntimeError("CAMADA APP DA ESTACAO INVALIDA.")
    if str(app) not in sys.path:
        sys.path.insert(0, str(app))
    os.environ["CJL_NETWORK_ROOT"] = str(master_root)
    os.environ["CJL_INSTALL_ROOT"] = str(Path(install_root).resolve())
    return project, master_root


def _json(value) -> None:
    print(json.dumps(value, ensure_ascii=False, indent=2, default=str))


def cmd_validate_master(root: Path) -> None:
    from Core.config import validate_deployment_root
    from Core.release import verify_manifest, verify_runtime_integrity
    validate_deployment_root(root)
    verify_manifest(root, exact_file_set=True)
    verify_runtime_integrity(root, exact_file_set=False, quick=True)
    _json({"ok": True, "master": str(root)})


COMPROMISED_PIN_SHA256 = {"4567f6e4059f4e6d795b79a102336e253421f7d965ee912fdd3485b3f4a88abd"}


def _pin_is_compromised(value: str) -> bool:
    return hashlib.sha256(str(value or "").strip().encode("utf-8")).hexdigest() in COMPROMISED_PIN_SHA256


class InvalidCredentialError(Exception):
    pass


class InsufficientAuthorityError(Exception):
    pass


def _master_admin_row(root: Path):
    # A credencial administrativa do bootstrap pertence ao Mestre. Ela nunca
    # pode ser validada contra o SQLite transitorio/cache da estacao.
    from Core.config import seed_database_path

    database = seed_database_path()
    if not database.is_file():
        raise RuntimeError(f"BANCO ADMINISTRATIVO DO MESTRE AUSENTE: {database}")
    uri = database.resolve().as_uri() + "?mode=ro"
    try:
        with sqlite3.connect(uri, uri=True, timeout=10) as connection:
            connection.row_factory = sqlite3.Row
            row = connection.execute(
                "SELECT nome,perfil,senha_hash,senha_salt,permissoes_json,ativo "
                "FROM usuarios WHERE nome='ADMIN' COLLATE NOCASE LIMIT 1"
            ).fetchone()
    except sqlite3.Error as exc:
        raise RuntimeError(f"BANCO ADMINISTRATIVO DO MESTRE NAO PODE SER LIDO: {exc}") from exc
    if row is None:
        raise RuntimeError("CONTA ADMIN PRINCIPAL AUSENTE NO BANCO DO MESTRE.")
    return row, database


def _verify_scrypt_pin(row, password: str) -> bool:
    if not row["senha_hash"] or not row["senha_salt"]:
        raise RuntimeError("CONTA ADMIN PRINCIPAL NAO POSSUI CREDENCIAL CONFIGURADA.")
    try:
        salt = base64.b64decode(str(row["senha_salt"]), validate=True)
        expected = base64.b64decode(str(row["senha_hash"]), validate=True)
        actual = hashlib.scrypt(str(password or "").encode("utf-8"), salt=salt, n=2**14, r=8, p=1, dklen=32)
    except (ValueError, TypeError) as exc:
        raise RuntimeError("CREDENCIAL ADMINISTRATIVA ARMAZENADA NO MESTRE E INVALIDA.") from exc
    return hmac.compare_digest(actual, expected)


def _system_admin_allowed(row) -> bool:
    if str(row["perfil"] or "").upper() != "ADMIN" or not bool(row["ativo"]):
        return False
    try:
        custom = json.loads(str(row["permissoes_json"] or "{}"))
    except ValueError:
        custom = {}
    if isinstance(custom, dict) and "SYSTEM_ADMIN" in custom:
        value = custom.get("SYSTEM_ADMIN")
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return bool(value)
        return str(value or "").strip().lower() in {"1", "true", "sim", "yes", "on"}
    return True


def cmd_validate_admin(root: Path) -> None:
    row, database = _master_admin_row(root)
    password = sys.stdin.readline().strip()
    if not _verify_scrypt_pin(row, password):
        raise InvalidCredentialError("CREDENCIAL ADMINISTRATIVA INVALIDA.")
    if not _system_admin_allowed(row):
        raise InsufficientAuthorityError("CONTA ADMIN PRINCIPAL SEM AUTORIDADE SYSTEM_ADMIN.")
    _json({
        "ok": True,
        "status": "ADMIN_OK",
        "administrator": str(row["nome"] or "ADMIN"),
        "database": str(database),
    })



def _windows_admin() -> bool:
    if os.name != "nt":
        return False
    try:
        import ctypes
        return bool(ctypes.windll.shell32.IsUserAnAdmin())
    except Exception:
        return False


def _hash_pin(pin: str) -> tuple[str, str]:
    value = str(pin or "").strip()
    if not (4 <= len(value) <= 32) or not value.isdigit():
        raise RuntimeError("NOVO PIN DEVE TER ENTRE 4 E 32 DIGITOS NUMERICOS.")
    salt = secrets.token_bytes(16)
    digest = hashlib.scrypt(value.encode("utf-8"), salt=salt, n=2**14, r=8, p=1, dklen=32)
    return base64.b64encode(digest).decode("ascii"), base64.b64encode(salt).decode("ascii")


def _current_security_id(root: Path) -> str:
    try:
        cfg=json.loads((root/"App"/"Config"/"sistema.json").read_text(encoding="utf-8"))
        sec=int((cfg.get("versioning") or {}).get("security") or 0)
        return f"SE-{sec:03d}"
    except Exception:
        return "SE-UNKNOWN"


def cmd_recover_admin(root: Path) -> None:
    if not _windows_admin():
        raise InsufficientAuthorityError("RECUPERACAO ADMIN EXIGE POWERSHELL/SYS_LOG EXECUTADO COMO ADMINISTRADOR DO WINDOWS.")
    from Core.config import seed_database_path
    from Core import stations

    active = stations.active()
    if active:
        raise RuntimeError("ENCERRE AS ESTACOES ATIVAS ANTES DA RECUPERACAO ADMINISTRATIVA.")
    first = sys.stdin.readline().strip()
    second = sys.stdin.readline().strip()
    if first != second:
        raise InvalidCredentialError("CONFIRMACAO DO NOVO PIN NAO CONFERE.")
    if _pin_is_compromised(first):
        raise InvalidCredentialError("NOVO PIN CONSTA NA LISTA LOCAL DE CREDENCIAIS LEGADAS COMPROMETIDAS.")
    password_hash, salt = _hash_pin(first)
    database = seed_database_path()
    if not database.is_file():
        raise RuntimeError("BANCO DO MESTRE AUSENTE.")
    now = datetime.now(ZoneInfo("America/Sao_Paulo")).isoformat(timespec="seconds")
    security_id = _current_security_id(root)
    with sqlite3.connect(database, timeout=15) as connection:
        connection.row_factory = sqlite3.Row
        rows = connection.execute("SELECT * FROM usuarios").fetchall()
        admin = next((row for row in rows if str(row["nome"] or "").upper() == "ADMIN"), None)
        if admin is None:
            raise RuntimeError("CONTA ADMIN PRINCIPAL AUSENTE. EXECUTE O BOOTSTRAP DO MESTRE ANTES DA RECUPERACAO.")
        for row in rows:
            if row["id"] != admin["id"] and row["senha_hash"] and row["senha_salt"] and _verify_scrypt_pin(row, first):
                raise InvalidCredentialError("NOVO PIN JA PERTENCE A OUTRO USUARIO.")
        connection.execute(
            "UPDATE usuarios SET perfil='ADMIN',senha_hash=?,senha_salt=?,ativo=1,trocar_senha=0,auth_version=auth_version+1,updated_at=? WHERE id=?",
            (password_hash, salt, now, admin["id"]),
        )
        connection.execute(
            "INSERT INTO app_meta(key,value) VALUES('admin_provisioning_required','0') ON CONFLICT(key) DO UPDATE SET value='0'"
        )
        connection.execute(
            "INSERT INTO app_meta(key,value) VALUES('security_admin_recovery_last',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (now,),
        )
        try:
            event_id = "AUD-" + secrets.token_hex(12).upper()
            connection.execute(
                "INSERT INTO auditoria_eventos(id,ocorrido_em,usuario_id,usuario_nome,estacao_id,evento,entidade_tipo,entidade_id,detalhes_json) VALUES(?,?,?,?,?,?,?,?,?)",
                (event_id, now, str(admin["id"]), "ADMIN", "LOCAL_WINDOWS_ADMIN", "ADMIN_PIN_RECOVERY", "USUARIO", str(admin["id"]), json.dumps({"authority":"WINDOWS_ADMIN_LOCAL","security":security_id})),
            )
        except sqlite3.Error:
            pass
    _json({"ok": True, "status": "ADMIN_RECOVERED", "administrator": "ADMIN", "database": str(database), "security": security_id})

def cmd_validate_admin_store(root: Path) -> None:
    row, database = _master_admin_row(root)
    if not _system_admin_allowed(row):
        raise RuntimeError("CONTA ADMIN PRINCIPAL NAO ESTA ATIVA COM AUTORIDADE SYSTEM_ADMIN.")
    if not row["senha_hash"] or not row["senha_salt"]:
        raise RuntimeError("CONTA ADMIN PRINCIPAL NAO POSSUI CREDENCIAL CONFIGURADA.")
    _json({"ok": True, "status": "ADMIN_STORE_OK", "administrator": str(row["nome"] or "ADMIN"), "database": str(database)})


def cmd_info(root: Path) -> None:
    from Core.release import release_state
    from Core.config import runtime_python, repository_root, shared_data_root, seed_database_path
    state = release_state(root)
    _json({"release": state, "master": str(root), "python": str(runtime_python()), "database": str(seed_database_path()), "repository": str(repository_root()), "shared_data": str(shared_data_root())})


def cmd_integrity(root: Path) -> None:
    from Core.release import verify_manifest, verify_runtime_integrity
    from Core.signature import verify_release_signature
    manifest = verify_manifest(root, exact_file_set=True)
    runtime = verify_runtime_integrity(root, exact_file_set=False, quick=True)
    lineage = verify_release_signature(root)
    _json({"ok": True, "app_integrity_files": len(manifest.get("files") or {}), "runtime_quick": True, "lineage": lineage})


def _db_integrity() -> dict:
    from Core.config import local_database_path, seed_database_path
    from Core.db import connect
    result = {}
    seed = seed_database_path()
    if seed.is_file():
        with sqlite3.connect(seed) as connection:
            result["seed"] = connection.execute("PRAGMA integrity_check").fetchone()[0]
    try:
        with connect() as connection:
            result["local"] = connection.execute("PRAGMA integrity_check").fetchone()[0]
    except Exception as exc:
        result["local"] = f"INDISPONIVEL: {exc}"
    result["local_path"] = str(local_database_path())
    result["seed_path"] = str(seed)
    return result


def cmd_diagnose(root: Path) -> None:
    from Core.config import repository_root, shared_data_root
    from Core.repository import repository_diagnostics
    data = {"database": _db_integrity(), "repository_path": str(repository_root()), "shared_data": str(shared_data_root())}
    try:
        data["repository"] = repository_diagnostics()
    except Exception as exc:
        data["repository"] = {"ok": False, "error": str(exc)}
    _json(data)


def cmd_stations(root: Path) -> None:
    from Core import stations
    _json({"stations": stations.active()})


def cmd_resources(root: Path) -> None:
    from Core import resources
    _json({"resources": resources.list_resources()})


def cmd_validate_installed(application_system: str, master: str, install_root: str) -> None:
    app_root, master_root = _configure_install(application_system, master, install_root)
    from Core.release import verify_manifest, verify_runtime_component
    from Core.config import runtime_python
    verify_manifest(app_root, exact_file_set=True)
    verify_runtime_component(app_root, "Python", runtime_root_override=Path(install_root).resolve() / "Runtime")
    python = runtime_python()
    if not python.is_file():
        raise RuntimeError("RUNTIME PYTHON LOCAL NAO FOI LOCALIZADO.")
    import sqlite3 as _sqlite3  # noqa: F401
    sys.path.insert(0, str(python.parent / "Lib" / "site-packages"))
    import openpyxl  # noqa: F401
    import PIL  # noqa: F401
    import tzdata  # noqa: F401
    _json({"ok": True, "application": str(app_root), "master": str(master_root), "python": str(python)})


def cmd_migration_precheck(root: Path) -> None:
    from Core import stations
    active = stations.active()
    if active:
        raise RuntimeError("EXISTEM ESTACOES ATIVAS. ENCERRAR/DESCONECTAR TODAS ANTES DA MIGRACAO.")
    db = _db_integrity()
    if db.get("seed") not in {None, "ok"}:
        raise RuntimeError(f"BANCO SEED FALHOU NO INTEGRITY_CHECK: {db.get('seed')}")
    _json({"ok": True, "stations": 0, "database": db})


def cmd_migration_postcheck(root: Path) -> None:
    from Core.config import repository_root, shared_data_root, seed_database_path, runtime_python
    required = [repository_root(), shared_data_root(), seed_database_path(), runtime_python()]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise RuntimeError("ESTRUTURA NOVA INCOMPLETA: " + "; ".join(missing))
    db = _db_integrity()
    if db.get("seed") != "ok":
        raise RuntimeError(f"BANCO FALHOU APOS MIGRACAO: {db.get('seed')}")
    _json({"ok": True, "database": db, "repository": str(repository_root()), "shared_data": str(shared_data_root()), "python": str(runtime_python())})


def main() -> int:
    if len(sys.argv) < 3:
        raise RuntimeError("COMANDO DO HOST NAO INFORMADO.")
    command = sys.argv[1].strip().lower()
    if command == "validate-installed":
        if len(sys.argv) < 5:
            raise RuntimeError("PARAMETROS DE INSTALACAO INCOMPLETOS.")
        cmd_validate_installed(sys.argv[2], sys.argv[3], sys.argv[4])
        return 0
    root = _master(sys.argv[2])
    commands = {
        "validate-master": cmd_validate_master,
        "validate-admin": cmd_validate_admin,
        "recover-admin": cmd_recover_admin,
        "validate-admin-store": cmd_validate_admin_store,
        "info": cmd_info,
        "integrity": cmd_integrity,
        "diagnose": cmd_diagnose,
        "stations": cmd_stations,
        "resources": cmd_resources,
        "migration-precheck": cmd_migration_precheck,
        "migration-postcheck": cmd_migration_postcheck,
        "database-check": lambda r: _json(_db_integrity()),
    }
    action = commands.get(command)
    if action is None:
        raise RuntimeError(f"COMANDO DO HOST NAO SUPORTADO: {command}.")
    try:
        action(root)
    except InvalidCredentialError as exc:
        _json({"ok": False, "status": "INVALID_CREDENTIAL", "message": str(exc)})
        return 10
    except InsufficientAuthorityError as exc:
        _json({"ok": False, "status": "INSUFFICIENT_AUTHORITY", "message": str(exc)})
        return 11
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except BaseException as exc:
        print(f"FALHA: {exc}", file=sys.stderr)
        raise SystemExit(1)
