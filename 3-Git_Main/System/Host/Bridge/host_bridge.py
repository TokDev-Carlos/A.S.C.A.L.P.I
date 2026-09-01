from __future__ import annotations

import base64
import gzip
import hashlib
import hmac
import json
import os
import secrets
import shutil
import sqlite3
import sys
import tempfile
import uuid
from contextlib import closing, contextmanager
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo


def _add_runtime_site_packages(root: Path) -> None:
    site_packages = root / "Runtime" / "Python" / "Lib" / "site-packages"
    if site_packages.is_dir():
        value = str(site_packages)
        if value not in sys.path:
            sys.path.insert(0, value)


def _master(value: str) -> Path:
    root = Path(value).expanduser().resolve()
    app = root / "App"
    _add_runtime_site_packages(root)
    if not (app / "Config" / "master.id").is_file():
        raise RuntimeError("MESTRE CJL System LAYOUT V4 INVALIDO.")
    if str(app) not in sys.path:
        sys.path.insert(0, str(app))
    os.environ["CJL_NETWORK_ROOT"] = str(root)
    return root


def _configure_install(application_app: str, master: str, install_root: str) -> tuple[Path, Path]:
    app = Path(application_app).resolve()
    project = app.parent
    master_root = Path(master).resolve()
    _add_runtime_site_packages(Path(install_root).resolve())
    if not (app / "Config" / "sistema.json").is_file():
        raise RuntimeError("CAMADA APP DA ESTACAO INVALIDA.")
    if str(app) not in sys.path:
        sys.path.insert(0, str(app))
    os.environ["CJL_NETWORK_ROOT"] = str(master_root)
    os.environ["CJL_INSTALL_ROOT"] = str(Path(install_root).resolve())
    return project, master_root


def _json(value) -> None:
    print(json.dumps(value, ensure_ascii=False, indent=2, default=str))


def cmd_validate_master(root: Path) -> None:
    from Core.config import validate_deployment_root
    from Core.release import verify_manifest, verify_runtime_integrity
    validate_deployment_root(root)
    verify_manifest(root, exact_file_set=True)
    verify_runtime_integrity(root, exact_file_set=False, quick=True)
    _json({"ok": True, "master": str(root)})


COMPROMISED_PIN_SHA256 = {"4567f6e4059f4e6d795b79a102336e253421f7d965ee912fdd3485b3f4a88abd"}


def _pin_is_compromised(value: str) -> bool:
    return hashlib.sha256(str(value or "").strip().encode("utf-8")).hexdigest() in COMPROMISED_PIN_SHA256


class InvalidCredentialError(Exception):
    pass


class InsufficientAuthorityError(Exception):
    pass


def _master_admin_row(root: Path):
    # A credencial administrativa do bootstrap pertence ao Mestre. Ela nunca
    # pode ser validada contra o SQLite transitorio/cache da estacao.
    from Core.config import seed_database_path

    database = seed_database_path()
    if not database.is_file():
        raise RuntimeError(f"BANCO ADMINISTRATIVO DO MESTRE AUSENTE: {database}")
    uri = database.resolve().as_uri() + "?mode=ro"
    try:
        with sqlite3.connect(uri, uri=True, timeout=10) as connection:
            connection.row_factory = sqlite3.Row
            row = connection.execute(
                "SELECT nome,perfil,senha_hash,senha_salt,permissoes_json,ativo "
                "FROM usuarios WHERE nome='ADMIN' COLLATE NOCASE LIMIT 1"
            ).fetchone()
    except sqlite3.Error as exc:
        raise RuntimeError(f"BANCO ADMINISTRATIVO DO MESTRE NAO PODE SER LIDO: {exc}") from exc
    if row is None:
        raise RuntimeError("CONTA ADMIN PRINCIPAL AUSENTE NO BANCO DO MESTRE.")
    return row, database


def _verify_scrypt_pin(row, password: str) -> bool:
    if not row["senha_hash"] or not row["senha_salt"]:
        raise RuntimeError("CONTA ADMIN PRINCIPAL NAO POSSUI CREDENCIAL CONFIGURADA.")
    try:
        salt = base64.b64decode(str(row["senha_salt"]), validate=True)
        expected = base64.b64decode(str(row["senha_hash"]), validate=True)
        actual = hashlib.scrypt(str(password or "").encode("utf-8"), salt=salt, n=2**14, r=8, p=1, dklen=32)
    except (ValueError, TypeError) as exc:
        raise RuntimeError("CREDENCIAL ADMINISTRATIVA ARMAZENADA NO MESTRE E INVALIDA.") from exc
    return hmac.compare_digest(actual, expected)


def _system_admin_allowed(row) -> bool:
    if str(row["perfil"] or "").upper() != "ADMIN" or not bool(row["ativo"]):
        return False
    try:
        custom = json.loads(str(row["permissoes_json"] or "{}"))
    except ValueError:
        custom = {}
    if isinstance(custom, dict) and "SYSTEM_ADMIN" in custom:
        value = custom.get("SYSTEM_ADMIN")
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return bool(value)
        return str(value or "").strip().lower() in {"1", "true", "sim", "yes", "on"}
    return True


def cmd_validate_admin(root: Path) -> None:
    row, database = _master_admin_row(root)
    password = sys.stdin.readline().strip()
    if not _verify_scrypt_pin(row, password):
        raise InvalidCredentialError("CREDENCIAL ADMINISTRATIVA INVALIDA.")
    if not _system_admin_allowed(row):
        raise InsufficientAuthorityError("CONTA ADMIN PRINCIPAL SEM AUTORIDADE SYSTEM_ADMIN.")
    _json({
        "ok": True,
        "status": "ADMIN_OK",
        "administrator": str(row["nome"] or "ADMIN"),
        "database": str(database),
    })



def _windows_admin() -> bool:
    if os.name != "nt":
        return False
    try:
        import ctypes
        return bool(ctypes.windll.shell32.IsUserAnAdmin())
    except Exception:
        return False


def _hash_pin(pin: str) -> tuple[str, str]:
    value = str(pin or "").strip()
    if not 4 <= len(value) <= 32:
        raise RuntimeError("NOVA CREDENCIAL DEVE TER ENTRE 4 E 32 CARACTERES.")
    salt = secrets.token_bytes(16)
    digest = hashlib.scrypt(value.encode("utf-8"), salt=salt, n=2**14, r=8, p=1, dklen=32)
    return base64.b64encode(digest).decode("ascii"), base64.b64encode(salt).decode("ascii")


def _current_security_id(root: Path) -> str:
    try:
        cfg=json.loads((root/"App"/"Config"/"sistema.json").read_text(encoding="utf-8"))
        sec=int((cfg.get("versioning") or {}).get("security") or 0)
        return f"SE-{sec:03d}"
    except Exception:
        return "SE-UNKNOWN"


def _now() -> str:
    return datetime.now(ZoneInfo("America/Sao_Paulo")).isoformat(timespec="seconds")


def _sha256(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _canonical_json(value: dict) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def _pretty_json(value: dict) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8")


def _read_required_json(path: Path, label: str) -> dict:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError) as exc:
        raise RuntimeError(f"{label} AUSENTE OU INVALIDO: {path}") from exc
    if not isinstance(value, dict):
        raise RuntimeError(f"{label} INVALIDO: {path}")
    return value


def _target_local_state(root: Path) -> Path:
    master_id = (root / "App" / "Config" / "master.id").read_text(encoding="utf-8-sig").strip().upper()
    normalized = str(root.resolve()).rstrip("\\/").upper()
    suffix = hashlib.sha256(normalized.encode("utf-8")).hexdigest()[:12].upper()
    local = Path(os.environ.get("LOCALAPPDATA") or (Path.home() / "AppData" / "Local"))
    return local / "CJL" / "Instancias" / f"{master_id}-{suffix}"


@contextmanager
def _recovery_repository_lock(root: Path, operation_id: str):
    lock = root / "Repo" / "Bloqueios" / "ESCRITA_GLOBAL.lock"
    lock.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "format": 1,
        "nonce": operation_id,
        "station_id": "LOCAL_WINDOWS_ADMIN",
        "pid": os.getpid(),
        "action": "ADMIN_CANONICAL_RECOVERY",
        "started_at": _now(),
    }
    try:
        descriptor = os.open(str(lock), os.O_WRONLY | os.O_CREAT | os.O_EXCL)
    except FileExistsError as exc:
        raise RuntimeError("O REPOSITORIO OFICIAL JA POSSUI UM BLOQUEIO DE ESCRITA.") from exc
    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(_pretty_json(payload))
            stream.flush()
            os.fsync(stream.fileno())
        yield
    finally:
        try:
            owner = _read_required_json(lock, "BLOQUEIO DE RECOVERY")
            if str(owner.get("nonce") or "") == operation_id:
                lock.unlink()
        except OSError:
            pass


def _migrate_schema_15(connection: sqlite3.Connection) -> None:
    columns = {str(row[1]) for row in connection.execute("PRAGMA table_info(planejamentos_rota)")}
    declarations = (
        ("tempo_estimado_min", "REAL NOT NULL DEFAULT 0"),
        ("motor_rota", "TEXT NOT NULL DEFAULT 'FALLBACK_LOCAL'"),
        ("distancia_fonte", "TEXT NOT NULL DEFAULT 'ESTIMADA'"),
        ("rota_geometria_json", "TEXT NOT NULL DEFAULT '[]'"),
        ("pedagios_json", "TEXT NOT NULL DEFAULT '[]'"),
        ("base_rodoviaria_versao", "TEXT NOT NULL DEFAULT ''"),
        ("base_pedagios_versao", "TEXT NOT NULL DEFAULT ''"),
    )
    for name, declaration in declarations:
        if name not in columns:
            connection.execute(f"ALTER TABLE planejamentos_rota ADD COLUMN {name} {declaration}")
    connection.execute(
        "INSERT INTO app_meta(key,value) VALUES('schema_version','15') "
        "ON CONFLICT(key) DO UPDATE SET value='15'"
    )


def _stage_admin_database(
    database: Path,
    password: str,
    password_hash: str,
    salt: str,
    now: str,
    operation_id: str,
    security_id: str,
    authority: str,
    expected_admin_id: str | None = None,
) -> dict:
    with closing(sqlite3.connect(database, timeout=15)) as connection:
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys=ON")
        connection.execute("PRAGMA journal_mode=DELETE")
        connection.execute("PRAGMA synchronous=FULL")
        connection.execute("BEGIN IMMEDIATE")
        _migrate_schema_15(connection)
        rows = connection.execute("SELECT * FROM usuarios").fetchall()
        admin = next((row for row in rows if str(row["nome"] or "").upper() == "ADMIN"), None)
        if admin is None:
            raise RuntimeError("CONTA ADMIN PRINCIPAL AUSENTE.")
        if expected_admin_id and str(admin["id"]) != expected_admin_id:
            raise RuntimeError("IDENTIDADE ADMIN DIVERGE ENTRE MESTRE E REPOSITORIO.")
        for row in rows:
            if row["id"] != admin["id"] and row["senha_hash"] and row["senha_salt"] and _verify_scrypt_pin(row, password):
                raise InvalidCredentialError("NOVA CREDENCIAL JA PERTENCE A OUTRO USUARIO.")
        connection.execute(
            "UPDATE usuarios SET perfil='ADMIN',senha_hash=?,senha_salt=?,ativo=1,trocar_senha=0,"
            "auth_version=COALESCE(auth_version,0)+1,updated_at=? WHERE id=?",
            (password_hash, salt, now, admin["id"]),
        )
        connection.execute(
            "INSERT INTO app_meta(key,value) VALUES('admin_provisioning_required','0') "
            "ON CONFLICT(key) DO UPDATE SET value='0'"
        )
        connection.execute(
            "INSERT INTO app_meta(key,value) VALUES('security_admin_recovery_last',?) "
            "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (now,),
        )
        try:
            connection.execute(
                "INSERT INTO auditoria_eventos(id,ocorrido_em,usuario_id,usuario_nome,estacao_id,evento,entidade_tipo,entidade_id,detalhes_json) "
                "VALUES(?,?,?,?,?,?,?,?,?)",
                (
                    "AUD-" + operation_id[:24].upper(), now, str(admin["id"]), "ADMIN",
                    "LOCAL_WINDOWS_ADMIN", "ADMIN_CANONICAL_RECOVERY", "USUARIO", str(admin["id"]),
                    json.dumps({"authority": authority, "security": security_id, "operation_id": operation_id}, ensure_ascii=False),
                ),
            )
        except sqlite3.Error:
            pass
        connection.commit()
        integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
        foreign_keys = connection.execute("PRAGMA foreign_key_check").fetchall()
        schema = connection.execute("SELECT value FROM app_meta WHERE key='schema_version'").fetchone()[0]
        if integrity != "ok" or foreign_keys or str(schema) != "15":
            raise RuntimeError("BANCO PREPARADO NAO PASSOU INTEGRIDADE/SCHEMA 15.")
        return {
            "id": str(admin["id"]),
            "name": str(admin["nome"]),
            "permissions": str(admin["permissoes_json"] or "{}"),
        }


def _recover_admin_canonical(
    root: Path,
    password: str,
    authority: str,
    *,
    expected_revision: int | None = None,
) -> dict:
    from Core import stations
    from Core.atomic import atomic_write_bytes, atomic_write_json
    from Core.config import repository_root, seed_database_path
    from Core.repository import _validate_head_chain

    active = stations.active()
    if active:
        raise RuntimeError("ENCERRE AS ESTACOES ATIVAS ANTES DA RECUPERACAO ADMINISTRATIVA.")
    if _pin_is_compromised(password):
        raise InvalidCredentialError("NOVA CREDENCIAL CONSTA NA LISTA LOCAL DE CREDENCIAIS LEGADAS COMPROMETIDAS.")
    password_hash, salt = _hash_pin(password)
    credential_fingerprint = _sha256((password_hash + ":" + salt).encode("ascii"))[:16]
    operation_id = uuid.uuid4().hex
    now = _now()
    security_id = _current_security_id(root)
    repo = repository_root()
    head_path = repo / "HEAD.json"
    master_database = seed_database_path()
    if not master_database.is_file():
        raise RuntimeError("BANCO DO MESTRE AUSENTE.")
    if master_database.with_name(master_database.name + "-wal").exists():
        with closing(sqlite3.connect(master_database, timeout=15)) as checkpoint:
            busy, _log, _checkpointed = checkpoint.execute("PRAGMA wal_checkpoint(TRUNCATE)").fetchone()
        if busy:
            raise RuntimeError("WAL DO MESTRE ESTA EM USO; FECHE O SISTEMA E REPITA.")
    head_before_bytes = head_path.read_bytes()
    head_before = _read_required_json(head_path, "HEAD OFICIAL")
    current_revision = int(head_before.get("revision") or 0)
    if expected_revision is not None and current_revision != expected_revision:
        raise RuntimeError(f"HEAD_CHANGED: esperado {expected_revision}, atual {current_revision}.")
    _validate_head_chain(head_before, {})
    current_snapshot = repo / "Revisoes" / f"{current_revision:09d}.sqlite.gz"
    current_transaction = repo / "Transacoes" / f"{current_revision:09d}.json"
    if _sha256(current_snapshot.read_bytes()) != str(head_before.get("snapshot_sha256") or "").lower():
        raise RuntimeError("HASH DO SNAPSHOT ATUAL DIVERGE DO HEAD.")
    current_transaction_hash = _sha256(current_transaction.read_bytes())
    if current_transaction_hash != str(head_before.get("transaction_file_sha256") or "").lower():
        raise RuntimeError("HASH DA TRANSACAO ATUAL DIVERGE DO HEAD.")

    revision = current_revision + 1
    revision_path = repo / "Revisoes" / f"{revision:09d}.sqlite.gz"
    transaction_path = repo / "Transacoes" / f"{revision:09d}.json"
    if revision_path.exists() or transaction_path.exists():
        raise RuntimeError(f"REVISAO DESTINO {revision} JA EXISTE.")
    rollback = root / "Logs" / "System" / "Recovery" / "Rollback" / operation_id
    rollback.mkdir(parents=True, exist_ok=False)
    shutil.copy2(master_database, rollback / "master-before.db")
    shutil.copy2(head_path, rollback / "HEAD-before.json")
    shutil.copy2(current_snapshot, rollback / current_snapshot.name)
    shutil.copy2(current_transaction, rollback / current_transaction.name)

    config = _read_required_json(root / "App" / "Config" / "sistema.json", "CONFIGURACAO DO SISTEMA")
    app_version = str(config.get("version") or "")
    schema_version = int(config.get("schema_version") or 0)
    if schema_version != 15:
        raise RuntimeError("RECOVERY CANONICO EXIGE SCHEMA 15.")
    local_state = _target_local_state(root)
    local_sync = local_state / "sync.json"
    local_sync_before = local_sync.read_bytes() if local_sync.is_file() else None
    if local_sync_before is not None:
        (rollback / "local-sync-before.json").write_bytes(local_sync_before)

    event_path = root / "Logs" / "System" / "Recovery" / f"admin-recovery-{operation_id}.json"
    master_before_bytes = master_database.read_bytes()
    committed = False
    with tempfile.TemporaryDirectory(prefix="cjl-admin-recovery-") as temporary:
        stage = Path(temporary)
        master_stage = stage / "master.sqlite"
        repo_stage = stage / "repo.sqlite"
        shutil.copy2(master_database, master_stage)
        repo_stage.write_bytes(gzip.decompress(current_snapshot.read_bytes()))
        master_admin = _stage_admin_database(
            master_stage, password, password_hash, salt, now, operation_id, security_id, authority
        )
        repo_admin = _stage_admin_database(
            repo_stage, password, password_hash, salt, now, operation_id, security_id, authority,
            expected_admin_id=master_admin["id"],
        )
        raw_repo = repo_stage.read_bytes()
        compressed = gzip.compress(raw_repo, compresslevel=6, mtime=0)
        snapshot_hash = _sha256(compressed)
        transaction = {
            "format": 2,
            "revision": revision,
            "previous_revision": current_revision,
            "previous_transaction_sha256": current_transaction_hash,
            "action": "ADMIN_CANONICAL_RECOVERY",
            "snapshot_sha256": snapshot_hash,
            "app_version": app_version,
            "schema_version": schema_version,
            "user_id": master_admin["id"],
            "user_name": "ADMIN",
            "station_id": "LOCAL_WINDOWS_ADMIN",
            "published_at": now,
            "operation_id": operation_id,
        }
        transaction["transaction_sha256"] = _sha256(_canonical_json(transaction))
        transaction_bytes = _pretty_json(transaction)
        transaction_file_hash = _sha256(transaction_bytes)
        head_after = {
            "format": 2,
            "revision": revision,
            "snapshot_sha256": snapshot_hash,
            "transaction_file_sha256": transaction_file_hash,
            "app_version": app_version,
            "schema_version": schema_version,
            "published_at": now,
            "station_id": "LOCAL_WINDOWS_ADMIN",
            "operation_id": operation_id,
        }
        try:
            with _recovery_repository_lock(root, operation_id):
                if head_path.read_bytes() != head_before_bytes:
                    raise RuntimeError("HEAD_CHANGED DURANTE O RECOVERY.")
                if revision_path.exists() or transaction_path.exists():
                    raise RuntimeError("REVISAO DESTINO SURGIU DURANTE O RECOVERY.")
                atomic_write_bytes(revision_path, compressed)
                atomic_write_bytes(transaction_path, transaction_bytes)
                atomic_write_bytes(master_database, master_stage.read_bytes())
                atomic_write_json(head_path, head_after)
                if local_sync.is_file():
                    local_sync.unlink()
                receipt = {
                    "format": 1,
                    "event": "ADMIN_CANONICAL_RECOVERY",
                    "status": "PASS",
                    "operation_id": operation_id,
                    "timestamp": now,
                    "timezone": "America/Sao_Paulo",
                    "authority": authority,
                    "security": security_id,
                    "administrator": "ADMIN",
                    "admin_id": master_admin["id"],
                    "permissions_preserved": master_admin["permissions"] == repo_admin["permissions"],
                    "credential_fingerprint": credential_fingerprint,
                    "head_before": current_revision,
                    "head_after": revision,
                    "master_sha256": _sha256(master_database.read_bytes()),
                    "snapshot_sha256": snapshot_hash,
                    "transaction_file_sha256": transaction_file_hash,
                    "local_resync_required": True,
                    "local_state": str(local_state),
                    "local_existed_before": local_state.exists(),
                    "rollback": str(rollback),
                }
                atomic_write_json(event_path, receipt)
                committed = True
        finally:
            if not committed:
                rollback_errors = []
                for action in (
                    lambda: atomic_write_bytes(master_database, master_before_bytes),
                    lambda: atomic_write_bytes(head_path, head_before_bytes),
                    lambda: revision_path.unlink(missing_ok=True),
                    lambda: transaction_path.unlink(missing_ok=True),
                    lambda: atomic_write_bytes(local_sync, local_sync_before) if local_sync_before is not None else None,
                ):
                    try:
                        action()
                    except OSError as exc:
                        rollback_errors.append(str(exc))
                if rollback_errors:
                    raise RuntimeError("ROLLBACK DO RECOVERY FALHOU: " + "; ".join(rollback_errors))
    return {
        "ok": True,
        "status": "ADMIN_RECOVERED_CANONICAL",
        "administrator": "ADMIN",
        "database": str(master_database),
        "security": security_id,
        "head_before": current_revision,
        "head_after": revision,
        "credential_fingerprint": credential_fingerprint,
        "event": str(event_path),
        "rollback": str(rollback),
        "local_resync_required": True,
    }


def cmd_recover_admin(root: Path) -> None:
    if not _windows_admin():
        raise InsufficientAuthorityError("RECUPERACAO ADMIN EXIGE POWERSHELL/SYS_LOG EXECUTADO COMO ADMINISTRADOR DO WINDOWS.")
    first = sys.stdin.readline().strip()
    second = sys.stdin.readline().strip()
    if first != second:
        raise InvalidCredentialError("CONFIRMACAO DA NOVA CREDENCIAL NAO CONFERE.")
    _json(_recover_admin_canonical(root, first, "WINDOWS_ADMIN_LOCAL"))


def cmd_admin_recovery_status(root: Path) -> None:
    try:
        row, database = _master_admin_row(root)
        with closing(sqlite3.connect(database, timeout=10)) as connection:
            meta = connection.execute(
                "SELECT value FROM app_meta WHERE key='admin_provisioning_required'"
            ).fetchone()
        provisioning = str(meta[0] if meta else "0").strip() == "1"
        required = provisioning or not bool(row["senha_hash"] and row["senha_salt"]) or not _system_admin_allowed(row)
        reason = "PROVISIONING_REQUIRED" if provisioning else ("ADMIN_STORE_INVALID" if required else "NORMAL")
    except Exception as exc:
        required = True
        reason = f"ADMIN_STORE_UNAVAILABLE: {exc}"
    _json({"ok": True, "recovery_required": required, "reason": reason})

def cmd_validate_admin_store(root: Path) -> None:
    row, database = _master_admin_row(root)
    if not _system_admin_allowed(row):
        raise RuntimeError("CONTA ADMIN PRINCIPAL NAO ESTA ATIVA COM AUTORIDADE SYSTEM_ADMIN.")
    if not row["senha_hash"] or not row["senha_salt"]:
        raise RuntimeError("CONTA ADMIN PRINCIPAL NAO POSSUI CREDENCIAL CONFIGURADA.")
    _json({"ok": True, "status": "ADMIN_STORE_OK", "administrator": str(row["nome"] or "ADMIN"), "database": str(database)})


def cmd_info(root: Path) -> None:
    from Core.release import release_state
    from Core.config import runtime_python, repository_root, shared_data_root, seed_database_path
    state = release_state(root)
    _json({"release": state, "master": str(root), "python": str(runtime_python()), "database": str(seed_database_path()), "repository": str(repository_root()), "shared_data": str(shared_data_root())})


def cmd_integrity(root: Path) -> None:
    from Core.release import verify_manifest, verify_runtime_integrity
    from Core.signature import verify_release_signature
    manifest = verify_manifest(root, exact_file_set=True)
    runtime = verify_runtime_integrity(root, exact_file_set=False, quick=True)
    lineage = verify_release_signature(root)
    _json({"ok": True, "app_integrity_files": len(manifest.get("files") or {}), "runtime_quick": True, "lineage": lineage})


def _db_integrity() -> dict:
    from Core.config import local_database_path, seed_database_path
    result = {}
    seed = seed_database_path()
    if seed.is_file():
        with sqlite3.connect(seed) as connection:
            result["seed"] = connection.execute("PRAGMA integrity_check").fetchone()[0]
    local = local_database_path()
    if local.is_file():
        try:
            with sqlite3.connect(local.resolve().as_uri() + "?mode=ro", uri=True, timeout=10) as connection:
                result["local"] = connection.execute("PRAGMA integrity_check").fetchone()[0]
        except Exception as exc:
            result["local"] = f"INDISPONIVEL: {exc}"
    else:
        result["local"] = "AUSENTE"
    result["local_path"] = str(local)
    result["seed_path"] = str(seed)
    return result


def cmd_diagnose(root: Path) -> None:
    from Core.config import repository_root, shared_data_root
    from Core.repository import status as repository_status
    data = {"database": _db_integrity(), "repository_path": str(repository_root()), "shared_data": str(shared_data_root())}
    try:
        data["repository"] = repository_status()
    except Exception as exc:
        data["repository"] = {"ok": False, "error": str(exc)}
    _json(data)


def cmd_stations(root: Path) -> None:
    from Core import stations
    _json({"stations": stations.active()})


def cmd_resources(root: Path) -> None:
    from Core import resources
    _json({"resources": resources.list_resources()})


def cmd_validate_installed(application_system: str, master: str, install_root: str) -> None:
    app_root, master_root = _configure_install(application_system, master, install_root)
    from Core.release import verify_manifest, verify_runtime_component
    from Core.config import runtime_python
    verify_manifest(app_root, exact_file_set=True)
    verify_runtime_component(app_root, "Python", runtime_root_override=Path(install_root).resolve() / "Runtime")
    python = runtime_python()
    if not python.is_file():
        raise RuntimeError("RUNTIME PYTHON LOCAL NAO FOI LOCALIZADO.")
    import sqlite3 as _sqlite3  # noqa: F401
    sys.path.insert(0, str(python.parent / "Lib" / "site-packages"))
    import openpyxl  # noqa: F401
    import PIL  # noqa: F401
    import tzdata  # noqa: F401
    _json({"ok": True, "application": str(app_root), "master": str(master_root), "python": str(python)})


def cmd_migration_precheck(root: Path) -> None:
    from Core import stations
    active = stations.active()
    if active:
        raise RuntimeError("EXISTEM ESTACOES ATIVAS. ENCERRAR/DESCONECTAR TODAS ANTES DA MIGRACAO.")
    db = _db_integrity()
    if db.get("seed") not in {None, "ok"}:
        raise RuntimeError(f"BANCO SEED FALHOU NO INTEGRITY_CHECK: {db.get('seed')}")
    _json({"ok": True, "stations": 0, "database": db})


def cmd_migration_postcheck(root: Path) -> None:
    from Core.config import repository_root, shared_data_root, seed_database_path, runtime_python
    required = [repository_root(), shared_data_root(), seed_database_path(), runtime_python()]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise RuntimeError("ESTRUTURA NOVA INCOMPLETA: " + "; ".join(missing))
    db = _db_integrity()
    if db.get("seed") != "ok":
        raise RuntimeError(f"BANCO FALHOU APOS MIGRACAO: {db.get('seed')}")
    _json({"ok": True, "database": db, "repository": str(repository_root()), "shared_data": str(shared_data_root()), "python": str(runtime_python())})


def main() -> int:
    if len(sys.argv) < 3:
        raise RuntimeError("COMANDO DO HOST NAO INFORMADO.")
    command = sys.argv[1].strip().lower()
    if command == "validate-installed":
        if len(sys.argv) < 5:
            raise RuntimeError("PARAMETROS DE INSTALACAO INCOMPLETOS.")
        cmd_validate_installed(sys.argv[2], sys.argv[3], sys.argv[4])
        return 0
    root = _master(sys.argv[2])
    commands = {
        "validate-master": cmd_validate_master,
        "validate-admin": cmd_validate_admin,
        "recover-admin": cmd_recover_admin,
        "admin-recovery-status": cmd_admin_recovery_status,
        "validate-admin-store": cmd_validate_admin_store,
        "info": cmd_info,
        "integrity": cmd_integrity,
        "diagnose": cmd_diagnose,
        "stations": cmd_stations,
        "resources": cmd_resources,
        "migration-precheck": cmd_migration_precheck,
        "migration-postcheck": cmd_migration_postcheck,
        "database-check": lambda r: _json(_db_integrity()),
    }
    action = commands.get(command)
    if action is None:
        raise RuntimeError(f"COMANDO DO HOST NAO SUPORTADO: {command}.")
    try:
        action(root)
    except InvalidCredentialError as exc:
        _json({"ok": False, "status": "INVALID_CREDENTIAL", "message": str(exc)})
        return 10
    except InsufficientAuthorityError as exc:
        _json({"ok": False, "status": "INSUFFICIENT_AUTHORITY", "message": str(exc)})
        return 11
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except BaseException as exc:
        print(f"FALHA: {exc}", file=sys.stderr)
        raise SystemExit(1)
